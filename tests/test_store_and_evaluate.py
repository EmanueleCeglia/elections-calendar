from datetime import date

from openpyxl import load_workbook

from elections_calendar import store
from elections_calendar.evaluate import date_verdict, evaluate
from elections_calendar.models import Change, Event, Period


def make_event(**kw):
    base = dict(
        event_id="DE-REG-saarland-2027", country_code="DE", country="Germany", level="regional",
        election_type="regional", indirect=False, region="Saarland", title="2027 Saarland state election",
        date_start=date(2027, 4, 18), date_end=date(2027, 4, 18), date_precision="day", date_text="18 April",
        status="confirmed", confidence="high", sources=["cor_pdf", "wikipedia"],
        source_urls=["https://a.example", "https://b.example"], first_seen=date(2026, 10, 12),
        last_updated=date(2026, 10, 12), last_seen=date(2026, 10, 12),
    )
    base.update(kw)
    return Event(**base)


def test_csv_roundtrip(tmp_path):
    path = tmp_path / "cal.csv"
    ev = make_event(second_round=date(2027, 5, 2), notes="x | y")
    store.write_calendar(path, [ev])
    back = store.read_calendar(path)[0]
    assert back.sources == ev.sources and back.source_urls == ev.source_urls
    assert back.second_round == date(2027, 5, 2)
    assert back.period == ev.period and back.first_seen == ev.first_seen


def test_changes_are_appended(tmp_path):
    path = tmp_path / "changes.csv"
    store.append_changes(path, [Change(date(2026, 10, 12), "a", "added")])
    store.append_changes(path, [Change(date(2026, 10, 19), "a", "refined", "date", "x", "y")])
    rows = store.read_changes(path)
    assert [r["change"] for r in rows] == ["added", "refined"]


def test_workbook_sheets(tmp_path):
    path = tmp_path / "cal.xlsx"
    store.write_workbook(path, [make_event()], [], [], [], [("Run date", "2026-10-12")])
    wb = load_workbook(path)
    assert wb.sheetnames == ["Calendar", "Past", "Needs review", "Changes", "About"]
    ws = wb["Calendar"]
    assert ws["A1"].value == "Date / from" and ws.max_row == 2
    assert ws.freeze_panes == "A2"


def test_date_verdicts():
    day = Period(date(2027, 4, 18), date(2027, 4, 18), "day")
    month = Period(date(2027, 4, 1), date(2027, 4, 30), "month")
    assert date_verdict(day, day) == "exact"
    assert date_verdict(day, month) == "consistent, less precise"
    assert date_verdict(month, day) == "consistent, more precise"
    assert date_verdict(day, Period(date(2027, 5, 2), date(2027, 5, 2), "day")) == "wrong day"


def test_evaluate(tmp_path):
    truth = tmp_path / "truth.csv"
    truth.write_text(
        "country_code,country,level,region,election_type,indirect,title,date_start,date_end,date_precision,second_round,status,date_text,notes,source\n"
        "DE,Germany,regional,Saarland,regional,no,2027 Saarland state election,2027-04-18,2027-04-18,day,,confirmed,18 April 2027,,\n"
        "DE,Germany,regional,Bremen,regional,no,2027 Bremen state election,2027-05-30,2027-05-30,day,,confirmed,30 May 2027,,\n"
    )
    cal = tmp_path / "cal.csv"
    store.write_calendar(cal, [make_event(), make_event(event_id="FR-NAT-x", country_code="FR", level="national",
                                                        election_type="presidential", region="")])
    r = evaluate(truth, cal, date(2026, 10, 12), date(2027, 12, 31))
    assert r["matched"] == 1 and r["recall"] == 0.5 and r["precision"] == 0.5
    assert [t.region for t in r["missing"]] == ["Bremen"]
    assert r["date_verdicts"] == {"exact": 1}
