"""Read/write the calendar.

The CSV is the pipeline's state: it is read back on the next run, and its
git diff shows what changed. The Excel workbook is regenerated from it on
every run for people to use (filters, frozen header, one sheet per view).
"""

from __future__ import annotations

import csv
from dataclasses import asdict
from datetime import date
from pathlib import Path

from .models import Change, Event

CALENDAR_COLUMNS = [
    "event_id", "country_code", "country", "level", "region", "election_type", "indirect", "title",
    "date_start", "date_end", "date_precision", "date_text", "second_round", "status", "confidence",
    "sources", "source_urls", "evidence", "notes", "wikidata_id", "wikipedia_title", "region_id",
    "aggregate", "first_seen", "last_updated", "last_seen", "seen_in_last_run",
]
CHANGE_COLUMNS = ["run_date", "event_id", "change", "field", "old", "new", "description"]
LIST_SEP = " | "

# Columns shown in the Excel "Calendar" sheet, with headers for people.
SHEET_COLUMNS = [
    ("date_start", "Date / from", 13), ("date_end", "To", 12), ("date_text", "As announced", 22),
    ("second_round", "Run-off", 12), ("country", "Country", 15), ("level", "Level", 10),
    ("region", "Region", 24), ("election_type", "Type", 16), ("indirect", "Indirect", 9),
    ("title", "Election", 46), ("status", "Status", 11), ("confidence", "Confidence", 11),
    ("date_precision", "Date precision", 13), ("sources", "Sources", 26), ("evidence", "Evidence", 60),
    ("notes", "Notes", 50), ("source_urls", "Links", 60), ("first_seen", "First seen", 12),
    ("last_updated", "Last updated", 12), ("event_id", "Event ID", 30),
]


def _to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, list):
        return LIST_SEP.join(value)
    return str(value)


def _date(value: str) -> date | None:
    return date.fromisoformat(value) if value else None


def event_to_row(ev: Event) -> dict[str, str]:
    data = asdict(ev)
    return {col: _to_str(data[col]) for col in CALENDAR_COLUMNS}


def row_to_event(row: dict[str, str]) -> Event:
    split = lambda v: [x for x in (v or "").split(LIST_SEP) if x]  # noqa: E731
    return Event(
        event_id=row["event_id"],
        country_code=row["country_code"],
        country=row["country"],
        level=row["level"],
        election_type=row["election_type"],
        indirect=row.get("indirect") == "yes",
        region=row.get("region", ""),
        title=row.get("title", ""),
        date_start=date.fromisoformat(row["date_start"]),
        date_end=date.fromisoformat(row["date_end"]),
        date_precision=row["date_precision"],
        date_text=row.get("date_text", ""),
        status=row.get("status", "estimated"),
        confidence=row.get("confidence", "medium"),
        second_round=_date(row.get("second_round", "")),
        region_id=row.get("region_id", ""),
        aggregate=row.get("aggregate") == "yes",
        sources=split(row.get("sources")),
        source_urls=split(row.get("source_urls")),
        wikidata_id=row.get("wikidata_id", ""),
        wikipedia_title=row.get("wikipedia_title", ""),
        evidence=row.get("evidence", ""),
        notes=row.get("notes", ""),
        first_seen=_date(row.get("first_seen", "")),
        last_updated=_date(row.get("last_updated", "")),
        last_seen=_date(row.get("last_seen", "")),
        seen_in_last_run=row.get("seen_in_last_run", "yes") == "yes",
    )


def sort_events(events: list[Event]) -> list[Event]:
    return sorted(events, key=lambda e: (e.date_start, e.date_end, e.country, e.region, e.title))


def read_calendar(path: Path) -> list[Event]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as fh:
        return [row_to_event(r) for r in csv.DictReader(fh)]


def write_calendar(path: Path, events: list[Event]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CALENDAR_COLUMNS)
        writer.writeheader()
        for ev in sort_events(events):
            writer.writerow(event_to_row(ev))


def read_changes(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def append_changes(path: Path, changes: list[Change]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    new_file = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CHANGE_COLUMNS)
        if new_file:
            writer.writeheader()
        for ch in changes:
            writer.writerow({k: _to_str(v) for k, v in asdict(ch).items()})


def write_review(path: Path, events: list[Event]) -> None:
    write_calendar(path, events)


def write_workbook(
    path: Path,
    upcoming: list[Event],
    past: list[Event],
    review: list[Event],
    changes: list[dict[str, str]],
    about: list[tuple[str, str]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    status_fill = {
        "confirmed": PatternFill("solid", fgColor="E2EFDA"),
        "expected": PatternFill("solid", fgColor="FFF2CC"),
        "estimated": PatternFill("solid", fgColor="FCE4D6"),
        "held": PatternFill("solid", fgColor="EDEDED"),
    }

    wb = Workbook()

    def add_sheet(title: str, columns, rows: list[list], status_col: int | None = None):
        ws = wb.create_sheet(title)
        ws.append([c[1] for c in columns])
        for i, (_, _, width) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for values in rows:
            ws.append(values)
            if status_col is not None:
                fill = status_fill.get(str(values[status_col]))
                if fill:
                    ws.cell(row=ws.max_row, column=status_col + 1).fill = fill
        for col_idx, (key, _, _) in enumerate(columns, start=1):
            if key in ("date_start", "date_end", "second_round", "first_seen", "last_updated", "run_date"):
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    row[0].number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        return ws

    def event_values(ev: Event) -> list:
        data = asdict(ev)
        out = []
        for key, _, _ in SHEET_COLUMNS:
            v = data[key]
            if isinstance(v, list):
                v = "\n".join(v) if key == "source_urls" else ", ".join(v)
            elif isinstance(v, bool):
                v = "yes" if v else "no"
            elif key in ("election_type", "level") and isinstance(v, str):
                v = v.replace("_", " ")
            out.append(v)
        return out

    status_idx = [k for k, _, _ in SHEET_COLUMNS].index("status")
    add_sheet("Calendar", SHEET_COLUMNS, [event_values(e) for e in sort_events(upcoming)], status_idx)
    add_sheet("Past", SHEET_COLUMNS, [event_values(e) for e in sort_events(past)], status_idx)
    add_sheet("Needs review", SHEET_COLUMNS, [event_values(e) for e in sort_events(review)], status_idx)
    change_cols = [("run_date", "Run date", 12), ("event_id", "Event ID", 30), ("change", "Change", 12),
                   ("field", "Field", 12), ("old", "Old", 34), ("new", "New", 34), ("description", "Description", 80)]
    change_rows = [
        [_date(c["run_date"]) if k == "run_date" else c.get(k, "") for k, _, _ in change_cols]
        for c in sorted(changes, key=lambda c: (c["run_date"], c["event_id"]), reverse=True)
    ]
    add_sheet("Changes", change_cols, change_rows)
    ws = add_sheet("About", [("key", "Item", 28), ("value", "Value", 110)], [list(r) for r in about])
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
